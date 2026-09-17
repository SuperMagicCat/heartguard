# -*- coding: utf-8 -*-
import glob
import json
from openpyxl import load_workbook

files = glob.glob(r'E:\heartguard\image\question\*.xlsx')
print('qfiles', files)
wb = load_workbook(files[0], data_only=True)
for name in wb.sheetnames:
    ws = wb[name]
    print('SHEET', name, 'rows', ws.max_row, 'cols', ws.max_column)
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i > 40:
            break
        print(i, row)
