# -*- coding: utf-8 -*-
import glob
import json
import re
from openpyxl import load_workbook

files = glob.glob(r'E:\heartguard\image\question\*.xlsx')
wb = load_workbook(files[0], data_only=True)
ws = wb[wb.sheetnames[0]]
prefix = re.compile(r'^[（(][A-Ca-cＡ-Ｃ][)）]\s*')

def clean(v):
    if v is None:
        return ''
    s = str(v).strip()
    return prefix.sub('', s)

bank = []
for i, row in enumerate(ws.iter_rows(values_only=True)):
    if i == 0:
        continue
    diff, num, q, ok, w1, w2 = row[:6]
    d = 1
    ds = str(diff or '')
    if '3' in ds:
        d = 3
    elif '2' in ds:
        d = 2
    bank.append({
        'd': d,
        'id': int(num) if num else i,
        'q': clean(q),
        'ok': clean(ok),
        'bad': [clean(w1), clean(w2)]
    })

out = r'E:\heartguard\rapid-bank.json'
with open(out, 'w', encoding='utf-8') as f:
    json.dump(bank, f, ensure_ascii=False, indent=2)
print('ok', bank[0]['ok'], '|', bank[0]['q'][:20])
print('count', len(bank))
